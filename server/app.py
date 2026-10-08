#!/usr/bin/env python3
"""
Lightweight REST API and Web Server for OmniBurn.
Built with Python standard library (http.server) for zero-dependency local execution.
"""

import http.server
import json
import os
import socketserver
import sys
import threading
import time
import urllib.parse
from datetime import datetime, timezone

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from engine.db import get_connection
from engine.sync_engine import sync_models, add_model, deprecate_model, restore_model, remove_model
from engine.recommender import recommend
from engine.telemetry_logger import log_task_run, get_telemetry_summary
from engine.toggles import toggle_model, toggle_provider
from engine.classifier import classify_workload
from engine.golden_harness import run_golden_suite, run_all_benchmarks

PORT = 8787
STATIC_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static")
ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def ensure_reconciled_workbook(target_path):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws_yields = wb.active
        ws_yields.title = "Model Yields & Economics"

        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        headers = [
            "Model ID", "Display Name", "Provider", "Tier", "Cost / Pool ($)",
            "Tasks / Pool Cycle", "Evidence", "API Val / Pool ($)", "Tasks / Month",
            "Cost / Task ($)", "API Cost / Task ($)", "Tokens / Done", "Sec / Done",
            "Quota Score", "Time Score", "Notes"
        ]
        ws_yields.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws_yields.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        conn = get_connection()
        cur = conn.cursor()
        rows = cur.execute("""
            SELECT y.model_id, m.display_name, m.provider, y.tier_id, y.cost_per_pool,
                   y.tasks_per_pool_cycle, y.tasks_per_pool_evidence, y.api_value_per_pool,
                   y.tasks_per_month, y.cost_per_completed_task, y.api_cost_per_task,
                   y.success_adjusted_tokens, y.success_adjusted_seconds, y.quota_first_score,
                   y.time_reliability_score, y.recommendation_notes
            FROM model_task_yields y
            JOIN models m ON y.model_id = m.model_id
            ORDER BY y.tier_id ASC, y.cost_per_completed_task ASC
        """).fetchall()

        for r in rows:
            ws_yields.append([
                r["model_id"], r["display_name"], r["provider"], f"Tier {r['tier_id']}",
                r["cost_per_pool"], r["tasks_per_pool_cycle"], r["tasks_per_pool_evidence"],
                r["api_value_per_pool"], r["tasks_per_month"], r["cost_per_completed_task"],
                r["api_cost_per_task"], r["success_adjusted_tokens"], r["success_adjusted_seconds"],
                r["quota_first_score"], r["time_reliability_score"], r["recommendation_notes"]
            ])

        for col in ws_yields.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_yields.column_dimensions[col_letter].width = min(50, max(max_len + 3, 12))

        ws_subs = wb.create_sheet("Subscriptions & Pools")
        ws_subs.append(["Subscription", "Provider", "Monthly Cost", "Pool Name", "Reset Window", "Type", "5h Quota %", "Weekly Quota %"])
        for col_idx in range(1, 9):
            c = ws_subs.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font

        sub_rows = cur.execute("""
            SELECT s.name, s.provider, s.monthly_cost, qp.pool_name, qp.reset_window_hours, qp.window_type, qp.current_pct_remaining, qp.weekly_pct_remaining
            FROM subscriptions s
            JOIN quota_pools qp ON s.id = qp.sub_id
        """).fetchall()
        for r in sub_rows:
            ws_subs.append(list(r))

        conn.close()
        wb.save(target_path)
        return target_path
    except Exception as e:
        print(f"Error generating workbook {target_path}: {e}")
        return None

class OmniBurnHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=STATIC_DIR, **kwargs)

    def _send_json(self, data, status_code=200):
        self.send_response(status_code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()
        self.wfile.write(json.dumps(data, indent=2).encode("utf-8"))

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        query = urllib.parse.parse_qs(parsed.query)

        # Static index.html fallback
        if path == "/" or path == "/index.html":
            index_path = os.path.join(STATIC_DIR, "index.html")
            if os.path.exists(index_path):
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.end_headers()
                with open(index_path, "rb") as f:
                    self.wfile.write(f.read())
                return

        # Direct file download support from root / Downloads
        if path.startswith("/download/"):
            fname = os.path.basename(path)
            cand_paths = [
                os.path.join(os.path.expanduser("~/Downloads"), fname),
                os.path.join(ROOT_DIR, fname)
            ]
            for cp in cand_paths:
                if os.path.exists(cp):
                    self.send_response(200)
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
                    self.end_headers()
                    with open(cp, "rb") as f:
                        self.wfile.write(f.read())
                    return
            
            # Dynamic fallback generation for .xlsx requests
            if fname.endswith(".xlsx"):
                target = os.path.join(ROOT_DIR, fname)
                generated = ensure_reconciled_workbook(target)
                if generated and os.path.exists(generated):
                    self.send_response(200)
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
                    self.end_headers()
                    with open(generated, "rb") as f:
                        self.wfile.write(f.read())
                    return

        # API Endpoints
        if path == "/api/status":
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM subscriptions")
            subs = [dict(r) for r in cursor.fetchall()]
            
            cursor.execute("SELECT * FROM quota_pools")
            pools = [dict(r) for r in cursor.fetchall()]

            cursor.execute("SELECT COUNT(*) as cnt FROM models WHERE is_active = 1")
            active_model_cnt = cursor.fetchone()["cnt"]

            cursor.execute("SELECT COUNT(*) as cnt FROM models")
            total_model_cnt = cursor.fetchone()["cnt"]

            cursor.execute("SELECT * FROM model_changelog ORDER BY id DESC LIMIT 5")
            recent_changes = [dict(r) for r in cursor.fetchall()]
            conn.close()

            active_cost = sum(s["monthly_cost"] for s in subs if s["status"] == "active")
            total_cost = sum(s["monthly_cost"] for s in subs)
            self._send_json({
                "status": "ok",
                "monthly_spend": round(active_cost, 2),
                "total_subscribed_cost": round(total_cost, 2),
                "daily_burn": round(active_cost / 30.416, 2),
                "subscriptions": subs,
                "quota_pools": pools,
                "active_models_count": active_model_cnt,
                "total_models_count": total_model_cnt,
                "recent_changelog": recent_changes
            })
            return

        elif path == "/api/recommend":
            tier = int(query.get("tier", [2])[0])
            strategy = query.get("strategy", ["quota_first"])[0]
            plan = query.get("plan", [None])[0]
            rec = recommend(tier_id=tier, strategy=strategy, plan_filter=plan)
            self._send_json(rec)
            return

        elif path == "/api/classify":
            prompt = query.get("prompt", [None])[0]
            repo_path = query.get("dir", [None])[0]
            cls = classify_workload(prompt=prompt, repo_path=repo_path)
            self._send_json(cls)
            return

        elif path == "/api/route":
            prompt = query.get("prompt", [None])[0]
            repo_path = query.get("dir", [None])[0]
            strat_override = query.get("strategy", [None])[0]
            plan = query.get("plan", [None])[0]
            cls = classify_workload(prompt=prompt, repo_path=repo_path)
            strat = strat_override or cls["recommended_strategy"]
            rec = recommend(tier_id=cls["tier_id"], strategy=strat, plan_filter=plan)
            self._send_json({
                "classification": cls,
                "recommendation": rec,
            })
            return

        elif path == "/api/benchmark":
            tier = query.get("tier", [None])[0]
            model = query.get("model", [None])[0]
            if tier:
                res = [run_golden_suite(int(tier), model_id=model)]
            else:
                res = run_all_benchmarks(model_id=model)
            self._send_json({"results": res})
            return

        elif path == "/api/models":
            conn = get_connection()
            cursor = conn.cursor()
            sort_key = query.get("sort", ["default"])[0]
            sort_dir = query.get("dir", ["asc"])[0].upper()
            if sort_dir not in ["ASC", "DESC"]:
                sort_dir = "ASC"

            order_clause = "ORDER BY y.tier_id ASC, m.is_active DESC, y.quota_first_score DESC"
            if sort_key == "sub_cost":
                order_clause = f"ORDER BY m.is_active DESC, y.cost_per_completed_task {sort_dir}"
            elif sort_key == "api_cost":
                order_clause = f"ORDER BY m.is_active DESC, y.api_cost_per_task {sort_dir}"
            elif sort_key == "tasks_month":
                order_clause = f"ORDER BY m.is_active DESC, y.tasks_per_month {sort_dir}"
            elif sort_key == "tasks_pool":
                order_clause = f"ORDER BY m.is_active DESC, y.tasks_per_pool_cycle {sort_dir}"
            elif sort_key == "cost_pool":
                order_clause = f"ORDER BY m.is_active DESC, y.cost_per_pool {sort_dir}"
            elif sort_key == "api_value":
                order_clause = f"ORDER BY m.is_active DESC, y.api_value_per_pool {sort_dir}"
            elif sort_key == "latency":
                order_clause = f"ORDER BY m.is_active DESC, y.success_adjusted_seconds {sort_dir}"
            elif sort_key == "tokens":
                order_clause = f"ORDER BY m.is_active DESC, y.success_adjusted_tokens {sort_dir}"
            elif sort_key == "name":
                order_clause = f"ORDER BY m.is_active DESC, m.display_name {sort_dir}"

            tier_param = query.get("tier", [None])[0]
            where_clause = ""
            params = []
            if tier_param:
                where_clause = "WHERE y.tier_id = ?"
                params.append(int(tier_param))

            cursor.execute(f"""
                SELECT m.model_id, m.display_name, m.provider, m.harness, m.pool_id,
                       m.reasoning_effort, m.speed_mode, m.is_frontier, m.is_active,
                       m.input_cost_per_m, m.output_cost_per_m,
                       y.tier_id, y.tasks_per_pool_cycle, y.tasks_per_month,
                       y.cost_per_completed_task, y.api_cost_per_task, y.cost_per_pool,
                       y.tasks_per_pool_low, y.tasks_per_pool_high, y.tasks_per_pool_evidence,
                       y.api_value_per_pool, y.leverage, y.success_adjusted_tokens,
                       y.success_adjusted_seconds, y.quota_first_score, y.time_reliability_score,
                       y.recommendation_notes
                FROM models m
                JOIN model_task_yields y ON m.model_id = y.model_id
                {where_clause}
                {order_clause}
            """, params)
            models = [dict(r) for r in cursor.fetchall()]
            conn.close()
            self._send_json({"models": models, "sorted_by": sort_key, "direction": sort_dir})
            return

        elif path == "/api/changelog":
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM model_changelog ORDER BY id DESC LIMIT 50")
            changes = [dict(r) for r in cursor.fetchall()]
            conn.close()
            self._send_json({"changelog": changes})
            return

        elif path == "/api/telemetry":
            summary = get_telemetry_summary()
            self._send_json({"telemetry_summary": summary})
            return

        # Default fallback
        super().do_GET()

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        length = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(length).decode("utf-8") if length > 0 else "{}"
        try:
            payload = json.loads(body)
        except Exception:
            payload = {}

        if path == "/api/sync":
            try:
                res = sync_models()
                self._send_json(res)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/classify":
            try:
                prompt = payload.get("prompt")
                repo_path = payload.get("dir")
                diff_text = payload.get("diff")
                cls = classify_workload(prompt=prompt, repo_path=repo_path, diff_text=diff_text)
                self._send_json(cls)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/route":
            try:
                prompt = payload.get("prompt")
                repo_path = payload.get("dir")
                diff_text = payload.get("diff")
                strat_override = payload.get("strategy")
                plan = payload.get("plan")
                cls = classify_workload(prompt=prompt, repo_path=repo_path, diff_text=diff_text)
                strat = strat_override or cls["recommended_strategy"]
                rec = recommend(tier_id=cls["tier_id"], strategy=strat, plan_filter=plan)
                self._send_json({
                    "classification": cls,
                    "recommendation": rec,
                })
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/toggle-model":
            try:
                model_id = payload.get("model_id")
                is_active = payload.get("is_active")  # None or bool
                if not model_id:
                    self._send_json({"error": "model_id required"}, status_code=400)
                    return
                res = toggle_model(model_id, is_active=is_active)
                self._send_json(res)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/models/add":
            try:
                res = add_model(payload)
                status = 400 if "error" in res else 200
                self._send_json(res, status_code=status)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/models/deprecate":
            try:
                mid = payload.get("model_id")
                reason = payload.get("reason", "Manually deprecated via API")
                res = deprecate_model(mid, reason=reason)
                status = 400 if "error" in res else 200
                self._send_json(res, status_code=status)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/models/restore":
            try:
                mid = payload.get("model_id")
                res = restore_model(mid)
                status = 400 if "error" in res else 200
                self._send_json(res, status_code=status)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/models/remove":
            try:
                mid = payload.get("model_id")
                res = remove_model(mid)
                status = 400 if "error" in res else 200
                self._send_json(res, status_code=status)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/toggle-provider":
            try:
                provider = payload.get("provider") or payload.get("sub_id")
                is_active = payload.get("is_active")
                if not provider:
                    self._send_json({"error": "provider or sub_id required"}, status_code=400)
                    return
                res = toggle_provider(provider, is_active=is_active)
                self._send_json(res)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/log-task":
            try:
                task_id = payload.get("task_id", f"RUN-{int(datetime.now().timestamp())}")
                model_id = payload.get("model_id", "gemini-3.8-flash-medium")
                task_class = payload.get("task_class", "Tier 2: Standard Engineering")
                completed = payload.get("completed", True)
                first_pass = payload.get("first_pass_success", True)
                attempts = payload.get("attempt_number", 1)
                tokens_in = payload.get("input_tokens", 6000)
                tokens_out = payload.get("output_tokens", 1500)
                sec = payload.get("wall_clock_seconds", 60.0)
                dec_5h = payload.get("dec_5h", 0)
                dec_wk = payload.get("dec_week", 0)
                desc = payload.get("description", "User interactive execution")

                res = log_task_run(
                    task_id=task_id, model_id=model_id, task_class=task_class,
                    completed=completed, first_pass_success=first_pass,
                    attempt_number=attempts, input_tokens=tokens_in,
                    output_tokens=tokens_out, wall_clock_seconds=sec,
                    dec_5h=dec_5h, dec_week=dec_wk, description=desc
                )
                self._send_json(res)
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        elif path == "/api/update-pool":
            try:
                pool_id = payload.get("pool_id")
                new_5h = payload.get("current_pct_remaining")
                new_wk = payload.get("weekly_pct_remaining")
                conn = get_connection()
                cursor = conn.cursor()
                cursor.execute("""
                    UPDATE quota_pools
                    SET current_pct_remaining = COALESCE(?, current_pct_remaining),
                        weekly_pct_remaining = COALESCE(?, weekly_pct_remaining)
                    WHERE id = ?
                """, (new_5h, new_wk, pool_id))
                conn.commit()
                conn.close()
                self._send_json({"status": "ok", "pool_id": pool_id})
            except Exception as e:
                self._send_json({"error": str(e)}, status_code=500)
            return

        self._send_json({"error": "Endpoint not found"}, status_code=404)

def start_background_sync(interval_hours=None):
    """
    Launches a background daemon thread that periodically runs sync_models.
    Defaults to 6 hours or SYNC_INTERVAL_HOURS environment variable.
    """
    if interval_hours is None:
        try:
            interval_hours = float(os.environ.get("SYNC_INTERVAL_HOURS", "6"))
        except ValueError:
            interval_hours = 6.0

    interval_seconds = max(300.0, interval_hours * 3600.0)

    def _sync_worker():
        time.sleep(15)  # initial delay before first background cycle
        while True:
            try:
                print(f"[OmniBurn] Running background periodic model sync...")
                res = sync_models()
                if res.get("changes_count", 0) > 0:
                    print(f"[OmniBurn] Background sync logged {res['changes_count']} changes.")
            except Exception as e:
                print(f"[OmniBurn] Background sync error: {e}")
            time.sleep(interval_seconds)

    thread = threading.Thread(target=_sync_worker, daemon=True, name="OmniBurnPeriodicSync")
    thread.start()
    return thread

def run_server(port=PORT):
    start_background_sync()
    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("", port), OmniBurnHandler) as httpd:
        print(f"OmniBurn server running at: http://localhost:{port}")
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\nShutting down server.")

if __name__ == "__main__":
    run_server()
